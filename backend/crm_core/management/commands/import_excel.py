import datetime
import os
import openpyxl
from django.core.management.base import BaseCommand
from crm_core.models import Client, MonthlyCalendar, CalendarItem, Employee, AdHocTask

REEL_URL = 'https://commondatastorage.googleapis.com/gtv-videos-bucket/sample/ForBiggerBlazes.mp4'
IMG_URL  = 'https://images.unsplash.com/photo-1579684385127-1ef15d508118?auto=format&fit=crop&w=800&q=80'

EMPLOYEES = [
    { "id": 1634, "user_id": "922614612", "first_name": "Shreya", "last_name": "Thakur", "email_id": "shreya.thakur@zappcode.in", "team_group": "Social Media Team", "mobile_number": "922614612", "designation": "Digital Marketing Manager", "avatar_color": "bg-indigo-600" },
    { "id": 1600, "user_id": "9834998359", "first_name": "Gaurav", "last_name": "Bhingare", "email_id": "gaurav.bhingare@zappcode.in", "team_group": "Social Media Team", "mobile_number": "9834998359", "designation": "Jr. Video Editor", "avatar_color": "bg-purple-600" },
    { "id": 1597, "user_id": "7447682328", "first_name": "Achal", "last_name": "Ragit", "email_id": "aachal.ragit@zappcode.in", "team_group": "Social Media Team", "mobile_number": "7447682328", "designation": "Graphic Designer", "avatar_color": "bg-pink-600" },
    { "id": 1617, "user_id": "9322621335", "first_name": "Rohan", "last_name": "Katre", "email_id": "rohankatre90@gmail.com", "team_group": "Social Media Team", "mobile_number": "9322621335", "designation": "Graphic Designer", "avatar_color": "bg-blue-600" },
    { "id": 1598, "user_id": "7517061696", "first_name": "Prachi", "last_name": "Pantawane", "email_id": "prachi.pantawane@zappcode.in", "team_group": "Social Media Team", "mobile_number": "7517061696", "designation": "Digital Marketing Executive", "avatar_color": "bg-teal-600" },
    { "id": 1594, "user_id": "8668988901", "first_name": "Achal", "last_name": "Wankhede", "email_id": "achal.wankhede@zappcode.in", "team_group": "Social Media Team", "mobile_number": "8668988901", "designation": "SEO Exceutive", "avatar_color": "bg-amber-600" },
    { "id": 1629, "user_id": "8805841001", "first_name": "Diksha", "last_name": "Bisen", "email_id": "bisendiksha97@gmail.com", "team_group": "Social Media Team", "mobile_number": "8805841001", "designation": "Intern Digital Marketing", "avatar_color": "bg-emerald-600" },
    { "id": 1623, "user_id": "9325053129", "first_name": "Prerna", "last_name": "Yewle", "email_id": "prernayewle091@gmail.com", "team_group": "Social Media Team", "mobile_number": "9325053129", "designation": "Intern eCommerce Management", "avatar_color": "bg-cyan-600" },
]

CONFIGS = [
    ('Neuro',              'Central Brain & Spine Hospital', 'neuro',        'NEURO',    '#7C3AED', 'Healthcare & Neurology',    20, 10, 30, False),
    ('nelson',             'Nelson',                        'nelson',        'NELSON',   '#2563EB', 'Healthcare & Maternity',    30, 20, 30, True ),
    ('OCC KCR',            'OCC KCR',                       'occ-kcr',       'OCC',      '#EA580C', 'Transport & Cab Services',  15,  8, 30, False),
    ('commerce',           'Prerna Commerce College',       'commerce',      'PRERNA',   '#059669', 'Education (Commerce)',       20, 10, 30, False),
    ('Kohinoor TEZ',       'Kohinoor TEZ',                  'kohinoor',      'KOHINOOR', '#D97706', 'Education (Study Notes)',    15, 10, 30, False),
    ('MBA',                'Prerna MBA',                    'mba',           'MBA',      '#9333EA', 'Higher Education',           20, 10, 30, False),
    ('Convent',            'Prerna Convent',                'convent',       'CONVENT',  '#0891B2', 'School Education',           20, 10, 30, False),
    ('international',      'Prerna International',          'international', 'INTL',     '#4F46E5', 'International Education',    20, 10, 30, False),
    ('zappkode solution ', 'Zappkode Solutions',            'zappkode',      'ZS',       '#1D4ED8', 'IT & Tech Agency',           20, 10, 30, False),
    ('zappcode academy',   'Zappkode Academy',              'zappacademy',   'ZA',       '#6D28D9', 'EdTech',                     20, 10, 30, False),
    ('Vibetara ',          'Vibetara',                      'vibetara',      'VIBE',     '#DB2777', 'Lifestyle & Fashion',        20, 10, 30, False),
]

def map_fmt(t):
    if not t: return 'IMAGE'
    u = str(t).upper()
    if 'REEL' in u: return 'REEL'
    if any(x in u for x in ['CARO', 'CROUS', 'CRUSE', 'SLIDE']): return 'CAROUSEL'
    return 'IMAGE'

def map_status(s):
    if not s: return 'READY_FOR_CLIENT'
    if 'DONE' in str(s).upper(): return 'APPROVED'
    return 'READY_FOR_CLIENT'

def fmt_date(d):
    if isinstance(d, datetime.datetime):
        return d.date()
    if isinstance(d, datetime.date):
        return d
    return None

class Command(BaseCommand):
    help = "Import October 2026 Excel Calendar and Employees into CRM"

    def handle(self, *args, **options):
        self.stdout.write("1. Seeding Employees...")
        emp_objs = {}
        for emp_info in EMPLOYEES:
            e, _ = Employee.objects.update_or_create(
                id=emp_info["id"],
                defaults={
                    "user_id": emp_info["user_id"],
                    "first_name": emp_info["first_name"],
                    "last_name": emp_info["last_name"],
                    "email_id": emp_info["email_id"],
                    "team_group": emp_info["team_group"],
                    "mobile_number": emp_info["mobile_number"],
                    "designation": emp_info["designation"],
                    "avatar_color": emp_info["avatar_color"],
                }
            )
            emp_objs[e.id] = e

        excel_path = r'F:\Oct Calender 2026.xlsx'
        if not os.path.exists(excel_path):
            self.stderr.write(f"Excel file not found at {excel_path}")
            return

        self.stdout.write(f"2. Loading Excel from {excel_path}...")
        wb = openpyxl.load_workbook(excel_path)

        for sheet_name, client_name, slug, code, color, industry, t_creatives, t_reels, t_stories, is_dual in CONFIGS:
            if sheet_name not in wb.sheetnames:
                self.stdout.write(f"Sheet {sheet_name} not found, skipping...")
                continue

            ws = wb[sheet_name]
            client, _ = Client.objects.update_or_create(
                slug=slug,
                defaults={
                    "name": client_name,
                    "code": code,
                    "primary_color": color,
                    "industry": industry,
                    "target_creatives": t_creatives,
                    "target_reels": t_reels,
                    "target_stories": t_stories,
                    "special_notes": f"Imported from {sheet_name} sheet.",
                }
            )

            calendar, _ = MonthlyCalendar.objects.update_or_create(
                client=client,
                month="October",
                year=2026,
                defaults={
                    "title": f"{client_name} - October 2026 Schedule",
                    "status": "SHARED",
                    "target_creatives": t_creatives,
                    "target_reels": t_reels,
                    "target_stories": t_stories,
                }
            )

            # Clear existing items for re-import
            calendar.items.all().delete()

            # Find data start row
            data_start = 20
            for r_idx, row in enumerate(ws.iter_rows(max_row=35, values_only=True), 1):
                if row and row[0] and 'sr' in str(row[0]).lower():
                    data_start = r_idx + 1
                    break

            time_choices = ['11:00 AM', '01:30 PM', '05:00 PM', '07:30 PM', '08:30 PM']
            emp_pool = [1600, 1597, 1617, 1598, 1629]
            sr_counter = 0

            if is_dual: # Nelson Dual Brand
                for row in ws.iter_rows(min_row=data_start, values_only=True):
                    if len(row) < 9: continue
                    d = fmt_date(row[1])
                    if not d or d.month != 10 or d.year != 2026: continue
                    day = str(row[2]) if row[2] else "Thursday"
                    special = str(row[3]) if row[3] else None

                    # Nelson Luxe:
                    l_p = str(row[4] or row[5] or "General").strip()
                    l_fmt = map_fmt(row[6])
                    l_v = str(row[7]) if len(row) > 7 and row[7] else None
                    l_c = str(row[8]) if len(row) > 8 and row[8] else None
                    l_s = map_status(row[9] if len(row) > 9 else None)

                    if row[6] or row[8]:
                        sr_counter += 1
                        CalendarItem.objects.create(
                            calendar=calendar,
                            sr_no=sr_counter,
                            scheduled_date=d,
                            scheduled_time=time_choices[sr_counter % len(time_choices)],
                            day_of_week=day,
                            special_day=special,
                            content_pillar=l_p,
                            format=l_fmt,
                            post_category='EXTRA' if sr_counter % 8 == 0 else 'REGULAR',
                            priority='HIGH' if special or sr_counter % 6 == 0 else 'NORMAL',
                            brand='Nelson Luxe Hospital',
                            visual_brief=l_v,
                            reel_bytes_script=l_c if l_fmt == 'REEL' else None,
                            caption=l_c if l_fmt != 'REEL' else None,
                            status=l_s,
                            media_type='reel' if l_fmt == 'REEL' else 'image',
                            media_url=REEL_URL if l_fmt == 'REEL' else IMG_URL,
                            assigned_to=emp_objs.get(1600 if l_fmt == 'REEL' else 1597)
                        )

                    # Nelson Hospital:
                    if len(row) > 11 and (row[11] or (len(row) > 15 and row[15])):
                        h_fmt = map_fmt(row[11])
                        h_p = str(row[12] or (row[13] if len(row) > 13 else None) or "General").strip()
                        h_v = str(row[14]) if len(row) > 14 and row[14] else None
                        h_c = str(row[15]) if len(row) > 15 and row[15] else None
                        h_s = map_status(row[16] if len(row) > 16 else None)

                        sr_counter += 1
                        CalendarItem.objects.create(
                            calendar=calendar,
                            sr_no=sr_counter,
                            scheduled_date=d,
                            scheduled_time="06:30 PM",
                            day_of_week=day,
                            special_day=special,
                            content_pillar=h_p,
                            format=h_fmt,
                            post_category='REGULAR',
                            priority='HIGH' if special else 'NORMAL',
                            brand='Nelson Hospital',
                            visual_brief=h_v,
                            reel_bytes_script=h_c if h_fmt == 'REEL' else None,
                            caption=h_c if h_fmt != 'REEL' else None,
                            status=h_s,
                            media_type='reel' if h_fmt == 'REEL' else 'image',
                            media_url=REEL_URL if h_fmt == 'REEL' else IMG_URL,
                            assigned_to=emp_objs.get(1617 if h_fmt != 'REEL' else 1600)
                        )
            else: # Standard sheets
                for row in ws.iter_rows(min_row=data_start, values_only=True):
                    if len(row) < 6: continue
                    d = fmt_date(row[1])
                    if not d or d.month != 10 or d.year != 2026: continue
                    day = str(row[2]) if row[2] else "Thursday"
                    special = str(row[3]) if row[3] else None
                    pillar = str(row[4] or "General").strip()
                    fmt = map_fmt(row[5])
                    content = str(row[6]) if len(row) > 6 and row[6] else None
                    visual = str(row[7]) if len(row) > 7 and row[7] else None
                    stat = map_status(row[8] if len(row) > 8 else None)

                    if not row[5] and not content: continue
                    sr_counter += 1

                    assigned_emp_id = 1600 if fmt == 'REEL' else emp_pool[sr_counter % len(emp_pool)]

                    CalendarItem.objects.create(
                        calendar=calendar,
                        sr_no=sr_counter,
                        scheduled_date=d,
                        scheduled_time=time_choices[sr_counter % len(time_choices)],
                        day_of_week=day,
                        special_day=special,
                        content_pillar=pillar,
                        format=fmt,
                        post_category='EXTRA' if sr_counter % 8 == 0 else 'REGULAR',
                        priority='HIGH' if special or sr_counter % 6 == 0 else ('MEDIUM' if sr_counter % 3 == 0 else 'NORMAL'),
                        visual_brief=visual,
                        reel_bytes_script=content if fmt == 'REEL' else None,
                        caption=content if fmt != 'REEL' else None,
                        status=stat,
                        media_type='reel' if fmt == 'REEL' else ('carousel' if fmt == 'CAROUSEL' else 'image'),
                        media_url=REEL_URL if fmt == 'REEL' else IMG_URL,
                        assigned_to=emp_objs.get(assigned_emp_id)
                    )

            self.stdout.write(f"[OK] Imported {client_name} ({sr_counter} items)")

        self.stdout.write("Excel import completed successfully!")
